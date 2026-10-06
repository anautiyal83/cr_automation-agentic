"""
Script to generate sample test fixture files for the cr_automation-agentic project.
Generates:
  - sample_ciq.xlsx  (CIQ Excel workbook with Site_Data and Device_Config sheets)
  - sample_mop.docx  (MOP document with CLI and REST API steps)
"""

import os
from pathlib import Path

import openpyxl
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

FIXTURES_DIR = Path(__file__).resolve().parent


def create_sample_ciq():
    """Create a minimal CIQ Excel file with two sheets."""
    wb = openpyxl.Workbook()

    # --- Sheet 1: Site_Data ---
    ws1 = wb.active
    ws1.title = "Site_Data"
    ws1.append(["Hostname", "IP_Address", "Subnet_Mask", "VLAN_ID", "Service_Code", "Location"])
    site_rows = [
        ["NY-BKN-001", "10.10.1.1",  "255.255.255.0", 100, "VL1001", "New York - Brooklyn"],
        ["NY-MHT-002", "10.10.2.1",  "255.255.255.0", 200, "VL1002", "New York - Manhattan"],
        ["TX-DAL-003", "10.20.1.1",  "255.255.255.0", 300, "VN2001", "Texas - Dallas"],
        ["CA-SFO-004", "10.30.1.1",  "255.255.255.0", 400, "VL1003", "California - San Francisco"],
        ["IL-CHI-005", "10.40.1.1",  "255.255.255.0", 500, "VN2002", "Illinois - Chicago"],
    ]
    for row in site_rows:
        ws1.append(row)

    # Auto-size columns for readability
    for col in ws1.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws1.column_dimensions[col[0].column_letter].width = max_len + 2

    # --- Sheet 2: Device_Config ---
    ws2 = wb.create_sheet(title="Device_Config")
    ws2.append(["Hostname", "Port", "Protocol", "Speed"])
    device_rows = [
        ["NY-BKN-001", "eth0", "BGP",  "1000Mbps"],
        ["TX-DAL-003", "eth1", "OSPF", "10000Mbps"],
        ["CA-SFO-004", "eth2", "ISIS", "10000Mbps"],
    ]
    for row in device_rows:
        ws2.append(row)

    for col in ws2.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws2.column_dimensions[col[0].column_letter].width = max_len + 2

    out_path = FIXTURES_DIR / "sample_ciq.xlsx"
    wb.save(out_path)
    print(f"Created {out_path}")


def create_sample_mop():
    """Create a minimal MOP Word document."""
    doc = Document()

    # Title
    title = doc.add_heading("VoLTE Provisioning MOP", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        "This Method of Procedure (MOP) describes the steps required to provision "
        "VoLTE services on the target network elements."
    )

    # --- Step 1: CLI ---
    doc.add_heading("Step 1 - Verify Interface Status", level=2)
    doc.add_paragraph("Execute the following CLI command to verify the interface is operational:")
    code1 = doc.add_paragraph()
    run1 = code1.add_run(
        "admin@router> show interface eth0"
    )
    run1.font.name = "Courier New"
    run1.font.size = Pt(9)
    doc.add_paragraph("Expected Output:")
    exp1 = doc.add_paragraph()
    run_exp1 = exp1.add_run(
        "Interface: eth0\n"
        "  Status: up\n"
        "  IP Address: 192.168.1.1/24\n"
        "  Speed: 1000Mbps\n"
        "  MTU: 1500"
    )
    run_exp1.font.name = "Courier New"
    run_exp1.font.size = Pt(9)

    # --- Step 2: CLI ---
    doc.add_heading("Step 2 - Configure IP Address", level=2)
    doc.add_paragraph("Enter configuration mode and assign the IP address:")
    code2 = doc.add_paragraph()
    run2 = code2.add_run(
        "admin@router> configure terminal\n"
        "router(config)# interface eth0\n"
        "router(config-if)# ip address {IP_Address} {Subnet_Mask}\n"
        "router(config-if)# no shutdown\n"
        "router(config-if)# exit"
    )
    run2.font.name = "Courier New"
    run2.font.size = Pt(9)
    doc.add_paragraph("Expected Output:")
    exp2 = doc.add_paragraph()
    run_exp2 = exp2.add_run("Configuration applied successfully.")
    run_exp2.font.name = "Courier New"
    run_exp2.font.size = Pt(9)

    # --- Step 3: CLI ---
    doc.add_heading("Step 3 - Create VLAN", level=2)
    doc.add_paragraph("Create the VLAN and assign a descriptive name:")
    code3 = doc.add_paragraph()
    run3 = code3.add_run(
        "router(config)# vlan {VLAN_ID}\n"
        "router(config-vlan)# name {Hostname}-VLAN\n"
        "router(config-vlan)# exit"
    )
    run3.font.name = "Courier New"
    run3.font.size = Pt(9)
    doc.add_paragraph("Expected Output:")
    exp3 = doc.add_paragraph()
    run_exp3 = exp3.add_run("VLAN {VLAN_ID} created successfully.")
    run_exp3.font.name = "Courier New"
    run_exp3.font.size = Pt(9)

    # --- Step 4: REST API ---
    doc.add_heading("Step 4 - Activate Service via REST API", level=2)
    doc.add_paragraph("Send a POST request to activate the VoLTE service on the provisioning platform:")
    code4 = doc.add_paragraph()
    run4 = code4.add_run(
        "POST /api/v1/services/activate\n"
        "Host: provision.example.com\n"
        "Content-Type: application/json\n"
        "Authorization: Bearer {token}\n"
        "\n"
        "{\n"
        '  "hostname": "{Hostname}",\n'
        '  "service_code": "{Service_Code}",\n'
        '  "vlan_id": {VLAN_ID},\n'
        '  "ip_address": "{IP_Address}"\n'
        "}"
    )
    run4.font.name = "Courier New"
    run4.font.size = Pt(9)
    doc.add_paragraph("Expected Response (HTTP 200):")
    exp4 = doc.add_paragraph()
    run_exp4 = exp4.add_run(
        '{\n'
        '  "status": "activated",\n'
        '  "service_id": "SVC-00123",\n'
        '  "message": "Service activated successfully"\n'
        '}'
    )
    run_exp4.font.name = "Courier New"
    run_exp4.font.size = Pt(9)

    out_path = FIXTURES_DIR / "sample_mop.docx"
    doc.save(out_path)
    print(f"Created {out_path}")


if __name__ == "__main__":
    create_sample_ciq()
    create_sample_mop()
    print("All fixture files generated successfully.")
