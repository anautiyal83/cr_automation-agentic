"""Test execution — invokes ciq-processor and cli-automation-standalone JARs.

Used in two modes:
- mode="ciq"  (Step 5): Test validation-rules.yaml + json-output.yaml via ciq-processor
- mode="exec" (Step 8): Test workflow.yaml + simulator-config.yaml via cli-automation-standalone
"""
from __future__ import annotations

import asyncio
import json
import logging
import re
import subprocess
from pathlib import Path
from typing import Any

from src.core.models import StepOutput, RetryContext

logger = logging.getLogger(__name__)


class TestExecutor:
    def __init__(self, mode: str = "ciq"):
        self.mode = mode  # "ciq" or "exec"
        if mode == "ciq":
            self.step_number = 5
            self.step_name = "CIQ Validation Tester"
        else:
            self.step_number = 8
            self.step_name = "Execution Tester"

    async def execute(
        self,
        input_data: dict,
        config: Any,
        few_shot_examples: list[dict] | None = None,
        retry_context: RetryContext | None = None,
    ) -> StepOutput:
        if self.mode == "ciq":
            return await self._run_ciq_test(input_data, config)
        else:
            return await self._run_exec_test(input_data, config)

    async def _run_ciq_test(self, input_data: dict, config: Any) -> StepOutput:
        """Test validation-rules.yaml and json-output.yaml via ciq-processor JAR.

        Runs multiple test cases:
        TC1: Validation-only (rules against CIQ, no JSON output)
        TC2: Full validation + JSON output generation
        TC3: JSON output structure validation (checks generated JSON)

        Then parses the validation report JSON for per-sheet details.
        """
        test_cases: list[dict[str, Any]] = []
        timeout = 300

        artifacts = input_data.get("artifacts", {})
        run_dir = input_data.get("run_dir", ".")
        uc_name = input_data.get("uc_name", "UC")

        node_type, activity = self._parse_uc_name(uc_name)

        # Save YAML artifacts to files for JAR consumption
        filename_map = {
            "validation_rules_yaml": f"{uc_name}_validation-rules.yaml",
            "json_template_yaml": f"{uc_name}_json-output.yaml",
        }
        for key in ("validation_rules_yaml", "json_template_yaml"):
            content = artifacts.get(key, "")
            if content:
                fname = filename_map.get(key, f"{uc_name}_{key.replace('_yaml', '')}.yaml")
                fpath = Path(run_dir) / fname
                fpath.write_text(content, encoding="utf-8")

        ciq_path = input_data.get("ciq_file_path", "")
        rules_path = str(Path(run_dir) / f"{uc_name}_validation-rules.yaml")
        json_output_path = str(Path(run_dir) / f"{uc_name}_json-output.yaml")
        has_json_config = Path(json_output_path).exists()

        ciq_jar = self._get_config(config, "ciq_processor_jar")

        if not (ciq_jar and Path(ciq_jar).exists()):
            logger.warning("ciq-processor JAR not found at: %s — skipping CIQ tests", ciq_jar)
            return StepOutput(
                status="failed",
                error={"test_failures": [{"test": "setup", "error": f"JAR not found: {ciq_jar}"}]},
                output_data={"test_cases": [], "artifacts": artifacts},
            )

        if not Path(rules_path).exists():
            return StepOutput(
                status="failed",
                error={"test_failures": [{"test": "setup", "error": "validation-rules.yaml not found"}]},
                output_data={"test_cases": [], "artifacts": artifacts},
            )

        base_cmd = [
            "java", "-jar", ciq_jar,
            "--mode", "ciq-validate",
            "--ciq", ciq_path,
            "--node-type", node_type,
            "--activity", activity,
            "--rules", rules_path,
        ]

        # ── TC1: Validation-only (no JSON output) ──
        tc1_dir = str(Path(run_dir) / "ciq_tc1_validation")
        Path(tc1_dir).mkdir(exist_ok=True)
        tc1 = await self._run_jar(
            "TC1: CIQ Validation (rules only)",
            base_cmd + ["--output", tc1_dir],
            timeout,
        )
        tc1["report"] = self._parse_validation_report(tc1_dir, uc_name)
        test_cases.append(tc1)

        # ── TC2: Full validation + JSON output ──
        if has_json_config:
            tc2_dir = str(Path(run_dir) / "ciq_tc2_full")
            Path(tc2_dir).mkdir(exist_ok=True)
            tc2 = await self._run_jar(
                "TC2: CIQ Validation + JSON Output Generation",
                base_cmd + [
                    "--output", tc2_dir,
                    "--json-output-config-file", json_output_path,
                    "--json-output-dir", tc2_dir,
                ],
                timeout,
            )
            tc2["report"] = self._parse_validation_report(tc2_dir, uc_name)
            test_cases.append(tc2)

            # ── TC3: JSON output structure check ──
            tc3 = self._check_json_output(tc2_dir, uc_name)
            test_cases.append(tc3)

        # Copy reports to main ciq_output dir for the pipeline
        main_output = str(Path(run_dir) / "ciq_output")
        Path(main_output).mkdir(exist_ok=True)
        src_dir = tc2_dir if has_json_config else tc1_dir
        for f in Path(src_dir).iterdir():
            dest = Path(main_output) / f.name
            if not dest.exists():
                dest.write_bytes(f.read_bytes())

        failures = [tc for tc in test_cases if tc["status"] == "FAILED"]
        status = "failed" if failures else "passed"

        # Generate consolidated test reports (HTML + Excel)
        html_path = str(Path(run_dir) / f"{uc_name}_ciq-test-report.html")
        excel_path = str(Path(run_dir) / f"{uc_name}_ciq-test-report.xlsx")
        self._generate_html_report(html_path, test_cases, uc_name, node_type, activity)
        self._generate_excel_report(excel_path, test_cases, uc_name, node_type, activity)

        return StepOutput(
            status=status,
            error={"test_failures": [{"test": tc["name"], "error": tc.get("error", "")} for tc in failures]} if failures else None,
            output_data={
                "test_cases": test_cases,
                "artifacts": artifacts,
                "report_html": html_path,
                "report_excel": excel_path,
            },
        )

    async def _run_jar(self, name: str, cmd: list[str], timeout: int) -> dict[str, Any]:
        """Run a JAR command and return structured test result."""
        try:
            result = await asyncio.to_thread(
                subprocess.run, cmd,
                capture_output=True, text=True, timeout=timeout,
            )
            return {
                "name": name,
                "status": "PASSED" if result.returncode == 0 else "FAILED",
                "returncode": result.returncode,
                "stdout": result.stdout,
                "stderr": result.stderr,
                "command": " ".join(cmd),
            }
        except subprocess.TimeoutExpired:
            return {"name": name, "status": "FAILED", "returncode": -1,
                    "error": f"Timeout after {timeout}s", "stdout": "", "stderr": ""}
        except Exception as e:
            return {"name": name, "status": "FAILED", "returncode": -1,
                    "error": str(e), "stdout": "", "stderr": ""}

    @staticmethod
    def _parse_validation_report(output_dir: str, uc_name: str) -> dict[str, Any]:
        """Parse the JSON validation report for detailed per-sheet results."""
        report_file = Path(output_dir) / f"{uc_name}_VALIDATION_REPORT.json"
        if not report_file.exists():
            return {}
        try:
            data = json.loads(report_file.read_text(encoding="utf-8"))
            return {
                "status": data.get("status", "UNKNOWN"),
                "totalErrors": data.get("totalErrors", 0),
                "globalErrors": data.get("globalErrors", []),
                "sheets": data.get("sheets", []),
                "parameters": data.get("parameters", {}),
            }
        except Exception:
            return {}

    @staticmethod
    def _check_json_output(output_dir: str, uc_name: str) -> dict[str, Any]:
        """TC3: Validate the generated JSON output file structure."""
        json_file = Path(output_dir) / f"{uc_name}.json"
        if not json_file.exists():
            return {
                "name": "TC3: JSON Output Structure Check",
                "status": "FAILED",
                "error": f"{uc_name}.json not generated",
                "details": {},
            }
        try:
            data = json.loads(json_file.read_text(encoding="utf-8"))
            cr_groups = data.get("crGroups", [])
            total_nodes = sum(
                len(ng.get("nodes", []))
                for cg in cr_groups
                for ng in cg.get("nodeGroups", [])
            )
            total_config_rows = sum(
                len(ng.get("configData", []))
                for cg in cr_groups
                for ng in cg.get("nodeGroups", [])
            )
            details = {
                "schemaVersion": data.get("schemaVersion"),
                "crGroups": len(cr_groups),
                "totalNodes": total_nodes,
                "totalConfigRows": total_config_rows,
                "hasNodeGroups": any(cg.get("nodeGroups") for cg in cr_groups),
            }
            # Basic validity checks
            errors = []
            if not cr_groups:
                errors.append("No crGroups in JSON output")
            if total_nodes == 0:
                errors.append("No nodes found in any nodeGroup")
            if total_config_rows == 0:
                errors.append("No configData rows found")

            return {
                "name": "TC3: JSON Output Structure Check",
                "status": "FAILED" if errors else "PASSED",
                "error": "; ".join(errors) if errors else "",
                "details": details,
            }
        except json.JSONDecodeError as e:
            return {
                "name": "TC3: JSON Output Structure Check",
                "status": "FAILED",
                "error": f"Invalid JSON: {e}",
                "details": {},
            }

    async def _run_exec_test(self, input_data: dict, config: Any) -> StepOutput:
        """Test workflow.yaml + simulator-config.yaml via cli-automation-standalone JAR."""
        failures: list[dict[str, Any]] = []
        results: dict[str, Any] = {}
        timeout = 300

        artifacts = input_data.get("artifacts", {})
        run_dir = input_data.get("run_dir", ".")
        uc_name = input_data.get("uc_name", "UC")

        # Save workflow and simulator config to files
        filename_map = {
            "workflow_yaml": f"{uc_name}.yaml",
            "simulator_config_yaml": f"{uc_name}_simulator-config.yaml",
        }
        for key in ("workflow_yaml", "simulator_config_yaml"):
            content = artifacts.get(key, "")
            if content:
                fname = filename_map.get(key, f"{uc_name}_{key.replace('_yaml', '')}.yaml")
                fpath = Path(run_dir) / fname
                fpath.write_text(content, encoding="utf-8")

        workflow_path = str(Path(run_dir) / f"{uc_name}.yaml")
        sim_config_path = str(Path(run_dir) / f"{uc_name}_simulator-config.yaml")

        cli_jar = self._get_config(config, "cli_automation_jar")

        if cli_jar and Path(cli_jar).exists() and Path(workflow_path).exists():
            try:
                cmd = ["java", "-jar", cli_jar,
                       "--workflow", workflow_path,
                       "--output", str(Path(run_dir) / "exec_result.json")]
                if Path(sim_config_path).exists():
                    cmd.extend(["--simulator-config", sim_config_path])

                result = await asyncio.to_thread(
                    subprocess.run, cmd,
                    capture_output=True, text=True, timeout=timeout,
                )
                results["execution_test"] = {
                    "returncode": result.returncode,
                    "stdout": result.stdout,
                    "stderr": result.stderr,
                }
                if result.returncode != 0:
                    failures.append({"test": "workflow_execution", "error": result.stderr or "Non-zero exit code"})
            except Exception as e:
                failures.append({"test": "workflow_execution", "error": str(e)})
        else:
            logger.warning("cli-automation JAR not found at: %s — skipping execution tests", cli_jar)

        status = "failed" if failures else "passed"
        return StepOutput(
            status=status,
            error={"test_failures": failures} if failures else None,
            output_data={"results": results, "artifacts": artifacts},
        )

    @staticmethod
    def _generate_html_report(
        filepath: str, test_cases: list[dict], uc_name: str,
        node_type: str, activity: str,
    ) -> None:
        """Generate a consolidated HTML test report."""
        from datetime import datetime, timezone
        import html as html_mod

        passed = sum(1 for tc in test_cases if tc.get("status") == "PASSED")
        failed = sum(1 for tc in test_cases if tc.get("status") == "FAILED")
        total = len(test_cases)
        overall = "PASSED" if failed == 0 and total > 0 else "FAILED"
        ts = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

        # Build test case rows
        tc_rows = ""
        for tc in test_cases:
            st = tc.get("status", "UNKNOWN")
            badge_cls = "pass" if st == "PASSED" else "fail"
            name = html_mod.escape(tc.get("name", ""))
            error = html_mod.escape(tc.get("error", ""))

            detail_parts = []

            # Validation report details
            report = tc.get("report", {})
            if report:
                sheets = report.get("sheets", [])
                if sheets:
                    detail_parts.append("<div class='sub-title'>Sheet Validation</div><table class='detail-tbl'>"
                                        "<tr><th>Sheet</th><th>Status</th><th>Rows</th><th>Errors</th></tr>")
                    for sh in sheets:
                        sh_st = sh.get("status", "?")
                        sh_cls = "pass" if sh_st == "PASSED" else "fail"
                        detail_parts.append(
                            f"<tr><td>{html_mod.escape(sh.get('sheetName',''))}</td>"
                            f"<td class='{sh_cls}'>{sh_st}</td>"
                            f"<td>{sh.get('rowsChecked',0)}</td>"
                            f"<td>{len(sh.get('errors',[]))}</td></tr>"
                        )
                        for err in sh.get("errors", [])[:10]:
                            msg = html_mod.escape(str(err.get("message", err)))
                            detail_parts.append(
                                f"<tr><td></td><td colspan='3' class='err-detail'>"
                                f"Row {err.get('row','?')}, {html_mod.escape(str(err.get('column','?')))}: {msg}</td></tr>"
                            )
                    detail_parts.append("</table>")

                global_errs = report.get("globalErrors", [])
                if global_errs:
                    detail_parts.append("<div class='sub-title'>Global Errors</div><ul class='err-list'>")
                    for ge in global_errs[:10]:
                        detail_parts.append(f"<li>{html_mod.escape(str(ge))}</li>")
                    detail_parts.append("</ul>")

                params = report.get("parameters", {})
                if params:
                    detail_parts.append("<div class='sub-title'>Output Parameters</div><table class='detail-tbl'>"
                                        "<tr><th>Parameter</th><th>Value</th></tr>")
                    for pk, pv in params.items():
                        detail_parts.append(f"<tr><td>{html_mod.escape(pk)}</td><td>{html_mod.escape(str(pv))}</td></tr>")
                    detail_parts.append("</table>")

            # JSON structure details (TC3)
            details = tc.get("details", {})
            if details:
                detail_parts.append("<div class='sub-title'>JSON Output Structure</div><table class='detail-tbl'>")
                for dk, dv in details.items():
                    detail_parts.append(f"<tr><td>{html_mod.escape(str(dk))}</td><td>{dv}</td></tr>")
                detail_parts.append("</table>")

            # Stdout/stderr
            stdout = tc.get("stdout", "").strip()
            stderr = tc.get("stderr", "").strip()
            if stdout:
                detail_parts.append(f"<div class='sub-title'>Output</div><pre>{html_mod.escape(stdout)}</pre>")
            if stderr:
                detail_parts.append(f"<div class='sub-title'>Errors</div><pre class='stderr'>{html_mod.escape(stderr)}</pre>")

            detail_html = "\n".join(detail_parts) if detail_parts else ""
            error_row = f"<div class='tc-error'>{error}</div>" if error else ""

            tc_rows += f"""
            <div class="tc-card">
              <div class="tc-header {badge_cls}-header">
                <span class="tc-badge badge-{badge_cls}">{st}</span>
                <span class="tc-name">{name}</span>
              </div>
              {error_row}
              <div class="tc-detail">{detail_html}</div>
            </div>"""

        html = f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>CIQ Test Report — {html_mod.escape(uc_name)}</title>
<style>
* {{ box-sizing: border-box; margin: 0; padding: 0; }}
body {{ font-family: 'Segoe UI', Arial, sans-serif; background: #f4f6f9; color: #333; font-size: 14px; }}
.header {{ background: #1a3a5c; color: #fff; padding: 20px 30px; }}
.header h1 {{ font-size: 22px; font-weight: 600; margin-bottom: 4px; }}
.header .meta {{ font-size: 13px; color: #b0c4d8; display: flex; gap: 24px; }}
.summary {{ display: flex; align-items: center; gap: 30px; padding: 18px 30px; margin: 16px 24px;
  border-radius: 8px; border-left: 6px solid; }}
.summary.pass {{ background: #eafaf1; border-color: #27ae60; }}
.summary.fail {{ background: #fdf2f2; border-color: #e74c3c; }}
.summary .status {{ font-size: 28px; font-weight: 700; }}
.summary.pass .status {{ color: #27ae60; }}
.summary.fail .status {{ color: #e74c3c; }}
.stats {{ display: flex; gap: 28px; }}
.stat {{ text-align: center; }}
.stat-num {{ font-size: 22px; font-weight: 700; color: #1a3a5c; }}
.stat-num.green {{ color: #27ae60; }}
.stat-num.red {{ color: #e74c3c; }}
.stat-label {{ font-size: 10px; text-transform: uppercase; color: #888; letter-spacing: 0.5px; }}
.tc-card {{ margin: 12px 24px; background: #fff; border-radius: 8px; box-shadow: 0 1px 3px rgba(0,0,0,0.08); overflow: hidden; }}
.tc-header {{ display: flex; align-items: center; gap: 12px; padding: 12px 18px; }}
.pass-header {{ background: #eafaf1; }}
.fail-header {{ background: #fdf2f2; }}
.tc-badge {{ font-size: 11px; font-weight: 700; padding: 3px 10px; border-radius: 10px; text-transform: uppercase; }}
.badge-pass {{ background: #27ae60; color: #fff; }}
.badge-fail {{ background: #e74c3c; color: #fff; }}
.tc-name {{ font-weight: 600; font-size: 14px; }}
.tc-error {{ padding: 8px 18px; color: #c0392b; font-size: 13px; background: #fef9f9; border-top: 1px solid #fce4e4; }}
.tc-detail {{ padding: 12px 18px; }}
.sub-title {{ font-size: 12px; font-weight: 600; color: #555; text-transform: uppercase; letter-spacing: 0.4px; margin: 10px 0 6px; }}
.sub-title:first-child {{ margin-top: 0; }}
.detail-tbl {{ width: 100%; border-collapse: collapse; font-size: 13px; margin-bottom: 8px; }}
.detail-tbl th {{ background: #f0f4f8; padding: 6px 12px; text-align: left; font-weight: 600; color: #555; border-bottom: 2px solid #dde3eb; }}
.detail-tbl td {{ padding: 5px 12px; border-bottom: 1px solid #eef0f3; }}
.detail-tbl .pass {{ color: #27ae60; font-weight: 600; }}
.detail-tbl .fail {{ color: #e74c3c; font-weight: 600; }}
.err-detail {{ color: #c0392b; font-size: 12px; font-style: italic; }}
.err-list {{ padding-left: 20px; color: #c0392b; font-size: 13px; }}
.err-list li {{ margin-bottom: 3px; }}
pre {{ background: #f8f9fa; padding: 10px 14px; border-radius: 4px; font-size: 12px; overflow-x: auto; white-space: pre-wrap; word-break: break-word; }}
pre.stderr {{ background: #fef5f5; color: #c0392b; }}
.footer {{ padding: 16px 30px; font-size: 11px; color: #999; text-align: center; margin-top: 12px; }}
</style></head>
<body>
<div class="header">
  <h1>CIQ Validation Test Report</h1>
  <div class="meta">
    <span><b>UC:</b> {html_mod.escape(uc_name)}</span>
    <span><b>Node Type:</b> {html_mod.escape(node_type)}</span>
    <span><b>Activity:</b> {html_mod.escape(activity)}</span>
    <span><b>Generated:</b> {ts}</span>
  </div>
</div>
<div class="summary {'pass' if overall == 'PASSED' else 'fail'}">
  <div class="status">{overall}</div>
  <div class="stats">
    <div class="stat"><div class="stat-num">{total}</div><div class="stat-label">Test Cases</div></div>
    <div class="stat"><div class="stat-num green">{passed}</div><div class="stat-label">Passed</div></div>
    <div class="stat"><div class="stat-num red">{failed}</div><div class="stat-label">Failed</div></div>
  </div>
</div>
{tc_rows}
<div class="footer">Generated by CR Automation Pipeline — Step 5: CIQ Validation Tester</div>
</body></html>"""

        Path(filepath).write_text(html, encoding="utf-8")

    @staticmethod
    def _generate_excel_report(
        filepath: str, test_cases: list[dict], uc_name: str,
        node_type: str, activity: str,
    ) -> None:
        """Generate a consolidated Excel test report."""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            logger.warning("openpyxl not available — skipping Excel report")
            return

        wb = Workbook()
        thin = Side(style="thin", color="CCCCCC")
        border = Border(bottom=thin)

        hdr_font = Font(bold=True, size=11, color="FFFFFF")
        hdr_fill = PatternFill("solid", fgColor="1A3A5C")
        pass_fill = PatternFill("solid", fgColor="EAFAF1")
        fail_fill = PatternFill("solid", fgColor="FDF2F2")
        pass_font = Font(bold=True, color="27AE60")
        fail_font = Font(bold=True, color="E74C3C")

        # ── Summary sheet ──
        ws = wb.active
        ws.title = "Test Summary"
        ws.append(["CIQ Validation Test Report"])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append(["UC Name", uc_name])
        ws.append(["Node Type", node_type])
        ws.append(["Activity", activity])
        ws.append([])

        headers = ["#", "Test Case", "Status", "Error"]
        ws.append(headers)
        for cell in ws[ws.max_row]:
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center")

        for i, tc in enumerate(test_cases, 1):
            st = tc.get("status", "UNKNOWN")
            row = [i, tc.get("name", ""), st, tc.get("error", "")]
            ws.append(row)
            r = ws.max_row
            fill = pass_fill if st == "PASSED" else fail_fill
            font = pass_font if st == "PASSED" else fail_font
            ws.cell(r, 3).font = font
            for c in range(1, 5):
                ws.cell(r, c).fill = fill
                ws.cell(r, c).border = border

        ws.column_dimensions["A"].width = 5
        ws.column_dimensions["B"].width = 45
        ws.column_dimensions["C"].width = 12
        ws.column_dimensions["D"].width = 60

        # ── Sheet validation details ──
        for tc in test_cases:
            report = tc.get("report", {})
            sheets = report.get("sheets", [])
            if not sheets:
                continue

            raw_title = tc.get("name", "Detail")
            # Excel sheet names cannot contain : \ / ? * [ ] and must be ≤31 chars
            safe_title = re.sub(r'[:\\/?\*\[\]]', '-', raw_title)[:31]
            ws2 = wb.create_sheet(title=safe_title)
            ws2.append(["Sheet", "Status", "Rows Checked", "Errors"])
            for cell in ws2[1]:
                cell.font = hdr_font
                cell.fill = hdr_fill
            for sh in sheets:
                sh_st = sh.get("status", "?")
                ws2.append([sh.get("sheetName", ""), sh_st, sh.get("rowsChecked", 0), len(sh.get("errors", []))])
                r = ws2.max_row
                ws2.cell(r, 2).font = pass_font if sh_st == "PASSED" else fail_font
                for err in sh.get("errors", [])[:20]:
                    ws2.append(["", f"Row {err.get('row','?')}, {err.get('column','?')}", "", str(err.get("message", err))])
                    ws2.cell(ws2.max_row, 4).font = Font(color="C0392B", italic=True, size=10)

            ws2.column_dimensions["A"].width = 25
            ws2.column_dimensions["B"].width = 15
            ws2.column_dimensions["C"].width = 14
            ws2.column_dimensions["D"].width = 60

            # Parameters
            params = report.get("parameters", {})
            if params:
                ws2.append([])
                ws2.append(["Output Parameters"])
                ws2.cell(ws2.max_row, 1).font = Font(bold=True, size=11)
                for pk, pv in params.items():
                    ws2.append([pk, str(pv)])

        # ── JSON structure details ──
        for tc in test_cases:
            details = tc.get("details", {})
            if not details:
                continue
            ws3 = wb.create_sheet(title="JSON Output")
            ws3.append(["Property", "Value"])
            for cell in ws3[1]:
                cell.font = hdr_font
                cell.fill = hdr_fill
            for dk, dv in details.items():
                ws3.append([str(dk), str(dv)])
            ws3.column_dimensions["A"].width = 25
            ws3.column_dimensions["B"].width = 30

        try:
            wb.save(filepath)
        except Exception as exc:
            logger.warning("Failed to save Excel report: %s", exc)

    @staticmethod
    def _parse_uc_name(uc_name: str) -> tuple[str, str]:
        """Extract node_type and activity from UC name.

        UC name format: <NODE_TYPE>_<ACTIVITY_NAME>
        e.g. DSR_1058_IFC_CAPABILITY_MODIFICATION → node_type=DSR, activity=1058_IFC_CAPABILITY_MODIFICATION
        """
        parts = uc_name.split("_", 1)
        if len(parts) == 2:
            return parts[0], parts[1]
        return uc_name, uc_name

    @staticmethod
    def _get_config(config: Any, key: str) -> str:
        if isinstance(config, dict):
            val = config.get(key, "")
        else:
            val = getattr(config, key, "")
        if val:
            p = Path(val)
            if not p.is_absolute():
                # Resolve relative paths from project root
                project_root = Path(__file__).resolve().parent.parent.parent
                p = project_root / p
            return str(p)
        return val
